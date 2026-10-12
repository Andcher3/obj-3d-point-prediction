from __future__ import annotations

import os
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(r"D:\Pycharm\post_graduate_project")
ASSET_DIR = ROOT / "reports" / "posters" / "assets"
POSTER_DIR = ROOT / "reports" / "posters"
WEEKLY_DIR = ROOT / "reports" / "weekly" / "2026-07-01-future-label-sequence"

ARCH = ASSET_DIR / "future-label-sequence-transformer-architecture.png"
TRAIN = ASSET_DIR / "future-label-sequence-training-pipeline.png"
SWITCH = ASSET_DIR / "future-label-sequence-switch-inference.png"

POSTER_PNG = POSTER_DIR / "2026-06-30-future-label-sequence-method-poster.png"
POSTER_PDF = POSTER_DIR / "2026-06-30-future-label-sequence-method-poster.pdf"
PPTX = WEEKLY_DIR / "2026-07-01-future-label-sequence-method.pptx"


INK = (20, 32, 51)
MUTED = (82, 97, 115)
PAPER = (251, 252, 253)
LINE = (214, 222, 232)
BLUE = (29, 78, 216)
BLUE_SOFT = (219, 234, 254)
TEAL = (15, 118, 110)
TEAL_SOFT = (204, 251, 241)
GREEN = (21, 128, 61)
GREEN_SOFT = (220, 252, 231)
AMBER = (180, 83, 9)
AMBER_SOFT = (254, 243, 199)
ROSE = (190, 18, 60)
ROSE_SOFT = (255, 228, 230)
SLATE_SOFT = (241, 245, 249)
WHITE = (255, 255, 255)


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\arial.ttf",
    ]
    for path in candidates:
        if os.path.exists(path):
            return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def wrap(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    lines: list[str] = []
    for para in text.split("\n"):
        if not para:
            lines.append("")
            continue
        cur = ""
        for ch in para:
            trial = cur + ch
            if draw.textlength(trial, font=fnt) <= max_width:
                cur = trial
            else:
                if cur:
                    lines.append(cur)
                cur = ch
        if cur:
            lines.append(cur)
    return lines


def draw_lines(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    lines: Iterable[str],
    fnt: ImageFont.FreeTypeFont,
    fill=INK,
    line_gap: int = 8,
) -> int:
    x, y = xy
    for line in lines:
        draw.text((x, y), line, font=fnt, fill=fill)
        bbox = draw.textbbox((x, y), line, font=fnt)
        y += bbox[3] - bbox[1] + line_gap
    return y


def rounded(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    fill,
    outline=LINE,
    radius: int = 14,
    width: int = 3,
) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def heading(draw: ImageDraw.ImageDraw, x: int, y: int, text: str, color=INK) -> int:
    draw.text((x, y), text, font=font(34, True), fill=color)
    return y + 48


def tag(draw: ImageDraw.ImageDraw, x: int, y: int, text: str, fill, color) -> int:
    fnt = font(21, True)
    w = int(draw.textlength(text, font=fnt)) + 30
    rounded(draw, (x, y, x + w, y + 40), fill=fill, outline=color, radius=8, width=2)
    draw.text((x + 15, y + 8), text, font=fnt, fill=color)
    return x + w + 10


def paste_fit(canvas: Image.Image, path: Path, box: tuple[int, int, int, int]) -> None:
    img = Image.open(path).convert("RGB")
    x1, y1, x2, y2 = box
    bw, bh = x2 - x1, y2 - y1
    scale = min(bw / img.width, bh / img.height)
    nw, nh = int(img.width * scale), int(img.height * scale)
    resized = img.resize((nw, nh), Image.Resampling.LANCZOS)
    ox, oy = x1 + (bw - nw) // 2, y1 + (bh - nh) // 2
    canvas.paste(resized, (ox, oy))


def bullet_block(draw: ImageDraw.ImageDraw, x: int, y: int, width: int, items: list[str], size: int = 24) -> int:
    fnt = font(size)
    for item in items:
        draw.ellipse((x, y + 10, x + 10, y + 20), fill=BLUE)
        lines = wrap(draw, item, fnt, width - 28)
        y = draw_lines(draw, (x + 24, y), lines, fnt, line_gap=8)
        y += 8
    return y


def make_poster() -> None:
    POSTER_DIR.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (3600, 2600), PAPER)
    d = ImageDraw.Draw(img)

    # Header
    d.text((70, 48), "方向一方法设想 | Smartphone IMU Activity Forecasting", font=font(28, True), fill=BLUE)
    d.text((70, 92), "基于 Future Patch Label Sequence 的动作延续与切换预测", font=font(62, True), fill=INK)
    subtitle = "预测当前与未来 patch 的标签序列，再由序列变化和分类置信度推导延续概率、切换时机和下一动作。"
    draw_lines(d, (72, 172), wrap(d, subtitle, font(26, True), 2050), font(26, True), fill=MUTED, line_gap=10)
    rounded(d, (2560, 56, 3515, 218), fill=TEAL_SOFT, outline=TEAL, radius=14, width=4)
    note = "组会定位：研究方案草案。已完成范式梳理、模型结构选择、训练策略设计和示意图制作；尚未宣称完成实验。"
    draw_lines(d, (2590, 84), wrap(d, note, font(25, True), 880), font(25, True), fill=INK, line_gap=10)
    d.line((70, 252, 3530, 252), fill=INK, width=5)

    # Layout boxes
    left_x, mid_x, right_x = 70, 910, 2550
    top = 290
    left_w, mid_w, right_w = 800, 1600, 980

    # Left: problem
    rounded(d, (left_x, top, left_x + left_w, top + 340), WHITE)
    y = heading(d, left_x + 24, top + 22, "1. 问题重述")
    d.text((left_x + 28, y), "输入：手机 IMU 连续序列", font=font(24, True), fill=BLUE)
    y += 42
    d.text((left_x + 28, y), "acc x/y/z + gyro x/y/z", font=font(25), fill=INK)
    y += 50
    d.text((left_x + 28, y), "输出：当前动作、未来动作序列、延续概率、切换时机、下一动作", font=font(23), fill=INK)
    y += 58
    rounded(d, (left_x + 28, y, left_x + left_w - 28, y + 98), fill=SLATE_SOFT, outline=LINE, radius=10)
    d.text((left_x + 48, y + 18), "past IMU patches -> p(y_t)...p(y_{t+H})", font=font(28, True), fill=INK)

    # Left: why not transition
    y0 = top + 365
    rounded(d, (left_x, y0, left_x + left_w, y0 + 500), WHITE)
    y = heading(d, left_x + 24, y0 + 22, "2. 为什么不主用 transition class")
    bullet_block(
        d,
        left_x + 28,
        y,
        left_w - 56,
        [
            "<sit-2-stand> 有直觉价值，但作为主标签会带来类别膨胀与样本稀疏。",
            "较新的相近任务更常预测未来标签序列或动作段，而非枚举所有 A-to-B 切换。",
            "未来序列中第一次 label 变化给出切换时机；变化后的 label 就是下一动作。",
        ],
        size=23,
    )
    rounded(d, (left_x + 28, y0 + 380, left_x + left_w - 28, y0 + 462), fill=GREEN_SOFT, outline=GREEN, radius=10)
    d.text((left_x + 52, y0 + 400), "transition type 可作为后续辅助头或分析标签。", font=font(24, True), fill=GREEN)

    # Left: related
    y0 = top + 890
    rounded(d, (left_x, y0, left_x + left_w, 2250), WHITE)
    y = heading(d, left_x + 24, y0 + 22, "3. 相关范式依据")
    related = [
        ("P2LHAP", "sensor patches -> patch-level label sequence，forecast 直接输出未来活动标签。"),
        ("FUTR", "learnable action queries 并行预测未来动作与 duration，避免自回归错误累积。"),
        ("ActFusion", "把未来不可见部分视作 mask tokens，通过补全动作序列完成 anticipation。"),
        ("本方法", "parallel query decoder 为主；自回归 decoder 作为后续消融。"),
    ]
    for title, desc in related:
        rounded(d, (left_x + 28, y, left_x + left_w - 28, y + 122), fill=(252, 253, 255), outline=LINE, radius=10)
        d.text((left_x + 48, y + 18), title, font=font(26, True), fill=BLUE)
        draw_lines(d, (left_x + 48, y + 56), wrap(d, desc, font(21), left_w - 110), font(21), fill=INK, line_gap=6)
        y += 142

    # Middle: architecture
    rounded(d, (mid_x, top, mid_x + mid_w, top + 880), WHITE)
    y = heading(d, mid_x + 24, top + 20, "4. 模型结构：Transformer Encoder + Parallel Query Decoder")
    paste_fit(img, ARCH, (mid_x + 20, y + 8, mid_x + mid_w - 20, top + 760))
    draw_lines(
        d,
        (mid_x + 36, top + 782),
        wrap(d, "历史 IMU patch tokens 进入 Transformer Encoder；q0...qH 表示当前与未来 horizon queries；decoder 并行输出 o0...oH，并由共享 MLP classifier 得到标签分布。", font(23, True), mid_w - 70),
        font(23, True),
        fill=INK,
        line_gap=8,
    )

    # Middle: training
    y0 = top + 905
    rounded(d, (mid_x, y0, mid_x + mid_w, 2250), WHITE)
    y = heading(d, mid_x + 24, y0 + 20, "5. 训练策略：无标签预训练 + 有标签微调")
    paste_fit(img, TRAIN, (mid_x + 20, y + 8, mid_x + mid_w - 20, y0 + 1080))
    rounded(d, (mid_x + 30, y0 + 1115, mid_x + mid_w - 30, y0 + 1228), fill=AMBER_SOFT, outline=AMBER, radius=10)
    correction = "注：Stage 2 的 label prediction 采用 unmasked parallel query decoder；mask 只用于 Stage 1 的 masked IMU reconstruction 预训练。"
    draw_lines(d, (mid_x + 55, y0 + 1142), wrap(d, correction, font(24, True), mid_w - 110), font(24, True), fill=(108, 59, 0), line_gap=8)

    # Right: switch inference
    rounded(d, (right_x, top, right_x + right_w, top + 715), WHITE)
    y = heading(d, right_x + 24, top + 20, "6. 如何推导切换输出")
    paste_fit(img, SWITCH, (right_x + 18, y + 4, right_x + right_w - 18, top + 565))
    rounded(d, (right_x + 24, top + 585, right_x + right_w - 24, top + 682), fill=SLATE_SOFT, outline=LINE, radius=10)
    d.text((right_x + 46, top + 602), "time_to_switch = first changed patch", font=font(24, True), fill=INK)
    d.text((right_x + 46, top + 637), "next_action = first changed label", font=font(24, True), fill=INK)

    # Right: losses
    y0 = top + 740
    rounded(d, (right_x, y0, right_x + right_w, y0 + 520), WHITE)
    y = heading(d, right_x + 24, y0 + 20, "7. 损失函数与样本构造")
    d.text((right_x + 30, y), "从连续流中取过去 L 个 patch 作为输入，未来 H 个 patch label 作为 target。", font=font(23), fill=INK)
    y += 64
    rounded(d, (right_x + 28, y, right_x + right_w - 28, y + 118), fill=SLATE_SOFT, outline=LINE, radius=10)
    d.text((right_x + 50, y + 18), "input  = X[t-L+1 : t]", font=font(27, True), fill=INK)
    d.text((right_x + 50, y + 62), "target = y[t : t+H]", font=font(27, True), fill=INK)
    y += 150
    rounded(d, (right_x + 28, y, right_x + right_w - 28, y + 118), fill=BLUE_SOFT, outline=BLUE, radius=10)
    d.text((right_x + 50, y + 18), "L = CE(y_t) + lambda * sum CE(y_{t+k})", font=font(25, True), fill=INK)
    d.text((right_x + 50, y + 64), "k = 1...H；远期 horizon 可降权", font=font(24), fill=INK)

    # Right: next plan
    y0 = top + 1285
    rounded(d, (right_x, y0, right_x + right_w, 2250), WHITE)
    y = heading(d, right_x + 24, y0 + 20, "8. 后续实验计划")
    plans = [
        ("数据", "连续 HAR 数据验证 future label sequence；UCI HAR 只做 encoder sanity check。", BLUE, BLUE_SOFT),
        ("结构", "early fusion、P2LHAP-style channel independent、hybrid cross-channel fusion。", TEAL, TEAL_SOFT),
        ("训练", "from scratch、masked reconstruction pretraining、future IMU regression。", GREEN, GREEN_SOFT),
        ("评价", "future accuracy、next-action accuracy、time-to-switch error、calibration。", AMBER, AMBER_SOFT),
    ]
    for title, desc, color, fill in plans:
        rounded(d, (right_x + 28, y, right_x + right_w - 28, y + 122), fill=fill, outline=color, radius=10)
        d.text((right_x + 50, y + 18), title, font=font(26, True), fill=color)
        draw_lines(d, (right_x + 138, y + 18), wrap(d, desc, font(21, True), right_w - 210), font(21, True), fill=INK, line_gap=7)
        y += 145
    rounded(d, (right_x + 28, y + 8, right_x + right_w - 28, y + 122), fill=ROSE_SOFT, outline=ROSE, radius=10)
    draw_lines(
        d,
        (right_x + 52, y + 32),
        wrap(d, "方向一扩展：从 single human label sequence 升级到 human_action_sequence + phone_motion_sequence。", font(23, True), right_w - 100),
        font(23, True),
        fill=ROSE,
        line_gap=7,
    )

    d.line((70, 2490, 3530, 2490), fill=INK, width=3)
    d.text((70, 2510), "Group meeting material for 2026-07-01 | Future Patch Label Sequence method proposal", font=font(22, True), fill=MUTED)
    d.text((2620, 2510), "Figures: imagegen | Poster/PPT: local generation", font=font(22, True), fill=MUTED)

    img.save(POSTER_PNG)
    img.save(POSTER_PDF, "PDF", resolution=220.0)


def add_textbox(slide, x, y, w, h, text, size=24, bold=False, color=INK, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    run = p.runs[0]
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    return shape


def add_box(slide, x, y, w, h, fill, line=LINE, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill)
    shape.line.color.rgb = RGBColor(*line)
    shape.line.width = Pt(1.25)
    return shape


def add_bullets(slide, x, y, w, h, title: str, bullets: list[str], accent=BLUE):
    add_box(slide, x, y, w, h, WHITE)
    add_textbox(slide, x + 0.22, y + 0.16, w - 0.44, 0.35, title, size=22, bold=True, color=accent)
    tx = slide.shapes.add_textbox(Inches(x + 0.28), Inches(y + 0.62), Inches(w - 0.5), Inches(h - 0.72))
    tf = tx.text_frame
    tf.clear()
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(17)
        p.font.color.rgb = RGBColor(*INK)
        p.space_after = Pt(7)


def slide_title(slide, title: str, subtitle: str | None = None):
    add_textbox(slide, 0.42, 0.22, 12.2, 0.48, title, size=28, bold=True, color=INK)
    if subtitle:
        add_textbox(slide, 0.45, 0.76, 12.0, 0.28, subtitle, size=14, bold=True, color=MUTED)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.08), Inches(13.333), Inches(0.04))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*BLUE)
    shape.line.color.rgb = RGBColor(*BLUE)


def make_pptx() -> None:
    WEEKLY_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Slide 1
    s = prs.slides.add_slide(blank)
    slide_title(s, "本周进展：从论文阅读到方法设想", "Future Patch Label Sequence for Smartphone IMU Activity Forecasting")
    add_box(s, 0.65, 1.45, 12.0, 1.05, BLUE_SOFT, BLUE)
    add_textbox(s, 0.92, 1.68, 11.4, 0.5, "问题转变：从单窗口 HAR 分类，转向连续 IMU 流上的当前识别 + 未来标签序列预测。", size=22, bold=True)
    add_bullets(s, 0.65, 2.85, 3.9, 3.9, "已有依据", ["P2LHAP：patch-to-label sequence", "FUTR：parallel action queries", "ActFusion：future mask completion"], accent=BLUE)
    add_bullets(s, 4.85, 2.85, 3.9, 3.9, "本方法选择", ["不用大量 transition class", "预测 p(y_t)...p(y_{t+H})", "由 label 序列推导切换"], accent=TEAL)
    add_bullets(s, 9.05, 2.85, 3.6, 3.9, "组会要点", ["方法结构图", "无标签预训练策略", "下一步实验计划"], accent=AMBER)

    # Slide 2
    s = prs.slides.add_slide(blank)
    slide_title(s, "任务定义：Future Patch Label Sequence", "输入过去 L 个 IMU patches，输出当前与未来 H 个 patch 的标签分布")
    add_textbox(s, 0.65, 1.45, 12.0, 0.45, "核心公式", size=20, bold=True, color=BLUE)
    add_box(s, 0.65, 1.95, 12.0, 1.0, SLATE_SOFT)
    add_textbox(s, 0.95, 2.2, 11.4, 0.45, "past IMU patches  ->  p(y_t), p(y_{t+1}), ..., p(y_{t+H})", size=24, bold=True)
    add_bullets(s, 0.65, 3.3, 5.7, 3.3, "为什么不主用 <A-2-B>", ["transition 类别数随动作组合增长", "切换样本少，容易类别不平衡", "序列预测能自然派生切换时机"], accent=ROSE)
    add_bullets(s, 6.7, 3.3, 5.95, 3.3, "输出解释", ["第一次 label 变化：time-to-switch", "变化后的 label：next action", "label confidence：continuation / switch probability"], accent=GREEN)

    # Slide 3
    s = prs.slides.add_slide(blank)
    slide_title(s, "模型结构：Parallel Query Transformer", "Transformer Encoder + Parallel Query Decoder + Shared MLP Classifier")
    s.shapes.add_picture(str(ARCH), Inches(0.45), Inches(1.25), width=Inches(12.45))
    add_box(s, 0.6, 6.72, 12.1, 0.46, AMBER_SOFT, AMBER)
    add_textbox(s, 0.83, 6.82, 11.7, 0.24, "q0 表示当前查询，q1...qH 表示未来 horizon；decoder 并行输出 o0...oH 后逐 token 分类。", size=13, bold=True, color=INK)

    # Slide 4
    s = prs.slides.add_slide(blank)
    slide_title(s, "训练策略：无标签预训练 + 有标签微调", "大量无监督 IMU 数据用于学习 dynamics，少量标签用于 future label prediction")
    s.shapes.add_picture(str(TRAIN), Inches(0.45), Inches(1.28), width=Inches(12.45))
    add_box(s, 0.55, 6.68, 12.2, 0.52, AMBER_SOFT, AMBER)
    add_textbox(s, 0.82, 6.79, 11.7, 0.28, "注意：label prediction 阶段采用 unmasked parallel query decoder；mask 只用于 Stage 1 的 masked reconstruction。", size=13, bold=True, color=(108, 59, 0))

    # Slide 5
    s = prs.slides.add_slide(blank)
    slide_title(s, "由未来标签概率推导切换时机与下一动作", "切换不是主标签，而是未来 label sequence 的结构性结果")
    s.shapes.add_picture(str(SWITCH), Inches(0.5), Inches(1.28), width=Inches(12.25))
    add_box(s, 0.7, 6.68, 5.95, 0.48, BLUE_SOFT, BLUE)
    add_textbox(s, 0.95, 6.79, 5.5, 0.24, "switch time = first changed patch", size=13, bold=True)
    add_box(s, 6.95, 6.68, 5.7, 0.48, GREEN_SOFT, GREEN)
    add_textbox(s, 7.2, 6.79, 5.2, 0.24, "next action = first changed label", size=13, bold=True)

    # Slide 6
    s = prs.slides.add_slide(blank)
    slide_title(s, "下一步实验计划", "先验证输出形式，再扩展到人体动作 + 手机/相机运动双标签")
    add_bullets(s, 0.65, 1.45, 3.9, 4.9, "数据与任务", ["UCI HAR：只做 encoder sanity check", "寻找连续 HAR 数据：验证 future labels", "后续自采：human action + phone motion"], accent=BLUE)
    add_bullets(s, 4.85, 1.45, 3.9, 4.9, "结构消融", ["early fusion", "P2LHAP-style channel independent", "hybrid stem + cross-channel fusion", "parallel decoder vs autoregressive decoder"], accent=TEAL)
    add_bullets(s, 9.05, 1.45, 3.6, 4.9, "评价指标", ["future label accuracy / macro-F1", "next-action accuracy", "time-to-switch error", "continuation probability calibration"], accent=AMBER)
    add_box(s, 0.65, 6.55, 12.0, 0.55, ROSE_SOFT, ROSE)
    add_textbox(s, 0.95, 6.68, 11.4, 0.25, "周三汇报结论：当前最合理的近期落地形式是 future patch label sequence，而不是枚举所有 transition class。", size=15, bold=True, color=ROSE)

    prs.save(PPTX)


def main() -> None:
    make_poster()
    make_pptx()
    print(f"saved poster PNG: {POSTER_PNG}")
    print(f"saved poster PDF: {POSTER_PDF}")
    print(f"saved PPTX: {PPTX}")


if __name__ == "__main__":
    main()
